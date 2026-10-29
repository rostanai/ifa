from docx import Document
from docx.shared import Pt

d = Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(11)


def h(t, n=1): d.add_heading(t, n)
def p(t): d.add_paragraph(t)
def bl(items):
    for i in items: d.add_paragraph(i, style="List Bullet")
def nl(items):
    for i in items: d.add_paragraph(i, style="List Number")


d.add_heading("Incident-to-Fix Agent", 0)
p("Nebius x NVIDIA Global AI Hackathon | Track: Coding & Agentic Engineering | Also targeting: Best Use of Tavily")
p("Tagline: A failing build goes in. A verified fix, with the commit that caused it, comes out.")

h("1. Problem")
p("When CI breaks, an engineer reads the log, finds the file, works out which commit caused it, searches for the "
  "error, writes a fix and re-runs the tests. This is slow, interrupts focus, and usually happens at the worst time. "
  "Most of it is mechanical work that an agent can do.")

h("2. Solution")
p("Incident-to-Fix Agent takes a failed CI log and a git repository. It pinpoints the failing code, finds the "
  "commit that introduced it, asks an NVIDIA Nemotron model (served on Nebius Token Factory) for the smallest root-cause "
  "fix, and re-runs the project's tests. It only commits the fix to a new branch, and optionally opens a pull request, "
  "if the tests pass. If they still fail, it reverts and reports honestly.")

h("3. How it works")
nl([
    "Parse: extract the innermost project stack frame and the error message (Python, JS/TS, Go, Ruby, Java).",
    "Blame: run git blame on the failing line to identify the culprit commit, author and message.",
    "Research (optional): query Tavily for documentation and known fixes for the exact error.",
    "Reason: send the error, culprit commit, file contents and docs to NVIDIA Nemotron on Nebius Token Factory; "
    "receive an explanation and the corrected file as JSON.",
    "Verify: apply the change on a fix/<commit> branch and run the repo's test command.",
    "Deliver: commit with a message that names the regression commit, and optionally push and open a PR with gh.",
])

h("4. Use of Nebius and NVIDIA (hackathon requirement)")
bl([
    "Inference runs on Nebius Token Factory through its OpenAI-compatible chat completions API.",
    "The model is an NVIDIA open-source model (default: Llama-3.1-Nemotron-Ultra-253B, configurable with NEBIUS_MODEL).",
    "Tavily provides live documentation lookup for the Tavily prize.",
])

h("5. Demo scenario")
p("The repository contains a function average(xs). Commit 1 handled an empty list. Commit 2 (\"perf: simplify "
  "average()\") removed that guard, so the test suite now fails with ZeroDivisionError. The agent reads the real CI "
  "log, identifies calc.py line 2 and the commit \"perf: simplify average()\", and proposes restoring the guard. "
  "It re-runs the tests, sees them pass, and commits the fix on a new branch.")

p("A second demo uses JavaScript. In user.js, displayName(u) originally guarded against a missing profile. The commit "
  "\"refactor: simplify displayName()\" removed the guard, so Node throws TypeError: Cannot read properties of "
  "undefined (reading 'name'). The agent reads the Node stack trace, takes the innermost frame, finds the refactor "
  "commit, adds optional chaining, and verifies the fix with node test.js. This shows the agent is not tied to one language.")

h("6. Why it can win (judging criteria)")
bl([
    "Technological Implementation: a closed loop of parse, blame, research, reason, verify and deliver, "
    "not a single prompt. A fix is never claimed without passing tests.",
    "Design: one command, readable step-by-step output, a clean commit and PR with an explanation.",
    "Potential Impact: every engineering team loses hours to broken builds; this removes the mechanical part.",
    "Quality of the Idea: it ties the fix to the regression commit, which makes the result explainable and reviewable.",
])

h("7. How to run")
bl([
    "export NEBIUS_API_KEY=...   (and optionally TAVILY_API_KEY, NEBIUS_MODEL)",
    "sh demo/make_demo.sh",
    "python3 agent.py demo/repo demo/ci.log --test 'python3 test_calc.py'",
    "JavaScript demo: sh demo/make_demo_js.sh, then python3 agent.py demo/repo-js demo/ci-js.log --test 'node test.js'",
    "Add --pr to push the branch and open a pull request with the gh CLI.",
    "python3 agent.py --selftest   (offline check of log parsing)",
])
p("The agent uses only the Python standard library and git. There is nothing to install.")

h("8. Current status and limitations")
bl([
    "Verified: log parsing selftest passes; the demo repo builds, fails as intended, and git blame resolves "
    "the culprit commit correctly.",
    "Verified end to end: the agent called NVIDIA Nemotron (nemotron-3.5-super-vl-120b-a12b), restored the empty-list "
    "guard, passed the tests and committed the fix on its own branch.",
    "Not yet verified on Nebius: that run used NVIDIA's hosted endpoint for development. The final demo must be "
    "re-run with NEBIUS_API_KEY against Nebius Token Factory, with the model ID copied from its model list.",
    "Whole-file rewrite is used instead of diffs for reliability; suited to small and medium files. Large files "
    "would need diff-based patching.",
    "Single failure per run. Multi-failure triage and a GitHub Actions trigger are future work.",
])

h("9. Roadmap")
bl([
    "GitHub Action that triggers on failed workflow runs and comments the proposed fix on the PR.",
    "Diff-based patching and multi-file fixes.",
    "A confidence score and a flaky-test detector before attempting a fix.",
])

h("10. Devpost submission checklist")
bl([
    "Project name and tagline (above).",
    "Description: sections 1 to 4 of this document.",
    "Built with: Python, Nebius Token Factory, NVIDIA Nemotron, Tavily, git, GitHub CLI.",
    "Public repository link.",
    "Demo video of about 3 minutes (script below).",
    "Select the track: Coding & Agentic Engineering.",
])

h("11. Three-minute demo script")
nl([
    "0:00 The problem: a red CI build at 5 pm. Show the failing log.",
    "0:30 Run the agent command. Narrate each numbered step as it prints.",
    "1:15 Show the culprit commit it found and the Nemotron explanation.",
    "1:45 Show the tests passing and the resulting commit diff.",
    "2:15 Show the Nebius Token Factory dashboard or request to prove where inference ran.",
    "2:40 Close on impact and the roadmap.",
])

d.save("/var/www/NVIDA/Incident-to-Fix-Agent.docx")
print("saved")
